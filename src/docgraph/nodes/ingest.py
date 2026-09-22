"""Carga de documentos multiformato."""

from __future__ import annotations

import logging
from pathlib import Path

from docgraph.models import Document
from docgraph.state import DocGraphState

logger = logging.getLogger(__name__)

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".yaml", ".yml", ".json"}


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    logger.info("%s: %d páginas", path.name, len(pages))
    return "\n".join(p for p in pages if p.strip())


def _read_docx(path: Path) -> str:
    from docx import Document as DocxDocument

    return "\n".join(p.text for p in DocxDocument(path).paragraphs)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


_READERS = {".pdf": _read_pdf, ".docx": _read_docx}


def read_document(path: Path) -> Document | None:
    reader = _READERS.get(path.suffix.lower(), _read_text)
    try:
        content = reader(path)
    except Exception as exc:  # noqa: BLE001 - un fichero ilegible no aborta el lote
        logger.error("No se pudo leer %s: %s", path.name, exc)
        return None

    if not content.strip():
        logger.warning("%s no contiene texto extraíble", path.name)
        return None

    return Document(path=path, content=content)


def ingest(state: DocGraphState) -> DocGraphState:
    input_dir: Path = state["input_dir"]

    if not input_dir.exists():
        return {"documents": [], "errors": [f"No existe el directorio {input_dir}"]}

    candidates = sorted(
        p for p in input_dir.glob("**/*")
        if p.is_file() and p.suffix.lower() in SUPPORTED
    )
    documents = [d for d in (read_document(p) for p in candidates) if d is not None]

    logger.info("Documentos cargados: %d de %d candidatos", len(documents), len(candidates))

    errors = []
    if not documents:
        errors.append(f"No se encontró ningún documento legible en {input_dir}")

    return {"documents": documents, "errors": errors}
